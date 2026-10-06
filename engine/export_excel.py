# -*- coding: utf-8 -*-
"""Export Excel de référence (openpyxl). Entrée : views.donnees_excel(plan)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def ecrire_excel(onglets, chemin, titre_projet=''):
    wb = Workbook(); wb.remove(wb.active)
    entete = PatternFill('solid', fgColor='1F4E78'); blanc = Font(bold=True, color='FFFFFF')
    fin = Side(style='thin', color='BFBFBF'); bord = Border(left=fin, right=fin, top=fin, bottom=fin)
    gris = PatternFill('solid', fgColor='F2F2F2')
    for o in onglets:
        ws = wb.create_sheet(o['onglet'])
        ws.append(o['colonnes'])
        for c in ws[1]:
            c.fill = entete; c.font = blanc; c.alignment = Alignment(horizontal='center', vertical='center'); c.border = bord
        for i, l in enumerate(o['lignes']):
            ws.append(l)
            for c in ws[ws.max_row]:
                c.border = bord
                if i % 2: c.fill = gris
                if l and 'libre' in l: c.font = Font(color='A6A6A6')
        for j, w in enumerate(o['largeurs'], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.oddHeader.center.text = titre_projet
    wb.save(chemin)
    return chemin
